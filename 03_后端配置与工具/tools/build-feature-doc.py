"""Build the editable feature introduction from docs/FEATURES.md.

Requires Python 3 and python-docx. Run from any directory; output stays in project.
The Markdown source is authoritative. No API credentials are read by this script.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs' / 'FEATURES.md'
OUTPUT = ROOT / '项目功能介绍.docx'


def font_style(style, size, bold=False):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')


def table(document, rows):
    grid = document.add_table(rows=0, cols=len(rows[0]))
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    widths = [1.12, 1.36, 4.22] if len(rows[0]) == 3 else [1.35, 5.35]
    if rows[0][0] == '文档版本':
        widths = [0.9, 1.25, 4.55]
    for col, width in zip(grid.columns, widths):
        col.width = Inches(width)
    props = grid._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        node = OxmlElement('w:' + edge)
        for name, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            node.set(qn('w:' + name), value)
        borders.append(node)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge in ['top', 'left', 'bottom', 'right']:
        node = OxmlElement('w:' + edge)
        node.set(qn('w:w'), '90')
        node.set(qn('w:type'), 'dxa')
        margins.append(node)
    props.append(margins)
    for i, row in enumerate(rows):
        cells = grid.add_row().cells
        trpr = grid.rows[-1]._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if i == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for j, value in enumerate(row):
            cells[j].width = Inches(widths[j])
            cells[j].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            r = p.add_run(value)
            r.font.size = Pt(10)
            r.bold = i == 0
            fill = OxmlElement('w:shd')
            fill.set(qn('w:fill'), 'E6EFF1' if i == 0 else 'FFFFFF')
            cells[j]._tc.get_or_add_tcPr().append(fill)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = 0.4
    spacer.add_run().font.size = Pt(2)


def main():
    doc = Document()
    # Remove inherited template rules, especially the built-in Title border.
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.9)
    font_style(doc.styles['Normal'], 10.5)
    normal = doc.styles['Normal'].paragraph_format
    normal.space_after = Pt(6)
    normal.line_spacing = 1.2
    for name, size in [('Title', 20), ('Heading 1', 15), ('Heading 2', 12)]:
        font_style(doc.styles[name], size, name != 'Title')
        doc.styles[name].paragraph_format.space_before = Pt(11 if name != 'Title' else 0)
        doc.styles[name].paragraph_format.space_after = Pt(6)
        doc.styles[name].paragraph_format.keep_with_next = True
    for name in ['List Bullet', 'List Number']:
        font_style(doc.styles[name], 10.5)
        doc.styles[name].paragraph_format.space_after = Pt(4)
        doc.styles[name].paragraph_format.line_spacing = 1.15
    doc.core_properties.title = '旅游拍照机位分享小程序功能介绍'
    doc.core_properties.subject = '功能说明与持续开发记录'
    doc.core_properties.author = 'photo-spot-share 项目组'
    doc.core_properties.keywords = '功能介绍, 微信小程序, 旅游, 机位分享'
    doc.core_properties.comments = ''
    lines = SOURCE.read_text(encoding='utf-8').splitlines()
    i = 0
    pending_page_break = False
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line == '<!-- pagebreak -->':
            # Break before the next heading, without an extra empty paragraph
            # that could overflow onto a blank page when the previous page is full.
            pending_page_break = True
            i += 1
            continue
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                values = [v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', v) for v in values):
                    rows.append(values)
                i += 1
            table(doc, rows)
            continue
        elif line.startswith('# '):
            doc.add_paragraph(line[2:], 'Title')
        elif line.startswith('## '):
            doc.add_paragraph(line[3:], 'Heading 1')
        elif line.startswith('### '):
            doc.add_paragraph(line[4:], 'Heading 2')
        elif line.startswith('- '):
            doc.add_paragraph(line[2:], 'List Bullet')
        elif re.match(r'^\d+\. ', line):
            doc.add_paragraph(re.sub(r'^\d+\. ', '', line), 'List Number')
        else:
            doc.add_paragraph(line)
        if pending_page_break:
            doc.paragraphs[-1].paragraph_format.page_break_before = True
            pending_page_break = False
        i += 1
    doc.save(OUTPUT)
    print(str(OUTPUT))


if __name__ == '__main__':
    main()
